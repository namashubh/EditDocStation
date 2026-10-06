"""All document / image processing operations."""
import base64
import io
import json
import os
import shutil
import subprocess
import tempfile
import threading
import zipfile
from pathlib import Path
from urllib.parse import urlparse

import pymupdf as fitz
from PIL import Image, ImageFilter, ImageOps


class ToolError(Exception):
    """User-facing error (bad input, wrong password, ...)."""


# ---------------------------------------------------------------- helpers

def parse_pages(spec, page_count):
    """'1-3, 5, 8-' -> [0, 1, 2, 4, 7, ...] (0-based, ordered, unique)."""
    spec = (spec or "").strip().lower()
    if spec in ("", "all"):
        return list(range(page_count))
    result = []
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                start = int(a) if a else 1
                end = int(b) if b else page_count
            else:
                start = end = int(part)
        except ValueError:
            raise ToolError(f"Invalid page range: '{part}'")
        if start < 1 or end > page_count or start > end:
            raise ToolError(f"Page range '{part}' is outside 1-{page_count}")
        for p in range(start - 1, end):
            if p not in result:
                result.append(p)
    if not result:
        raise ToolError("No pages selected")
    return result


def parse_range_groups(spec, page_count):
    """'1-3, 4-6' -> [[0,1,2],[3,4,5]] (one group per output file)."""
    groups = [parse_pages(part, page_count) for part in (spec or "").split(",") if part.strip()]
    if not groups:
        raise ToolError("Please enter at least one range, e.g. 1-3, 4-6")
    return groups


def hex_to_rgb(value, default=(0, 0, 0)):
    value = (value or "").lstrip("#")
    if len(value) != 6:
        return default
    try:
        return tuple(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except ValueError:
        return default


def open_pdf(path, password=""):
    try:
        doc = fitz.open(path)
    except Exception:
        raise ToolError(f"'{Path(path).name}' is not a valid PDF file")
    if doc.needs_pass and not doc.authenticate(password or ""):
        doc.close()
        raise ToolError(f"'{Path(path).name}' is password protected. Use 'Unlock PDF' first.")
    return doc


def zip_files(paths, out_zip):
    with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in paths:
            zf.write(p, Path(p).name)
    return out_zip


def single_or_zip(paths, work, zip_name):
    paths = [Path(p) for p in paths]
    if len(paths) == 1:
        return paths[0], paths[0].name
    return zip_files(paths, work / zip_name), zip_name


def clean_stem(path):
    # Uploads are saved as "001_name.ext"; drop the numeric prefix.
    stem = Path(path).stem
    return stem.split("_", 1)[1] if "_" in stem and stem[:3].isdigit() else stem


# ---------------------------------------------------------------- organize

def merge_pdfs(files, out):
    if len(files) < 2:
        raise ToolError("Select at least 2 PDF files to merge")
    result = fitz.open()
    for f in files:
        with open_pdf(f) as src:
            result.insert_pdf(src)
    result.save(out, garbage=3, deflate=True)
    return out


def split_pdf(src, work, mode, ranges):
    doc = open_pdf(src)
    stem = clean_stem(src)
    n = doc.page_count
    groups = [[i] for i in range(n)] if mode == "all" else parse_range_groups(ranges, n)
    outputs = []
    for group in groups:
        part = fitz.open()
        for p in group:
            part.insert_pdf(doc, from_page=p, to_page=p)
        label = f"{group[0] + 1}" if len(group) == 1 else f"{group[0] + 1}-{group[-1] + 1}"
        out = work / f"{stem}_page_{label}.pdf"
        part.save(out, garbage=3, deflate=True)
        outputs.append(out)
    return outputs


def remove_pages(src, out, pages):
    doc = open_pdf(src)
    to_delete = parse_pages(pages, doc.page_count)
    if len(to_delete) >= doc.page_count:
        raise ToolError("You cannot remove all pages")
    doc.delete_pages(sorted(to_delete))
    doc.save(out, garbage=3, deflate=True)
    return out


def rotate_pdf(src, out, angle, pages):
    doc = open_pdf(src)
    for p in parse_pages(pages, doc.page_count):
        page = doc[p]
        page.set_rotation((page.rotation + angle) % 360)
    doc.save(out, garbage=3, deflate=True)
    return out


# ---------------------------------------------------------------- optimize

COMPRESS_LEVELS = {
    "low": (150, 80),       # less compression, best quality
    "recommended": (110, 60),
    "extreme": (72, 35),    # strongest compression
}


def compress_pdf(src, out, level):
    dpi, quality = COMPRESS_LEVELS.get(level, COMPRESS_LEVELS["recommended"])
    doc = open_pdf(src)
    try:
        doc.rewrite_images(dpi_threshold=dpi + 10, dpi_target=dpi, quality=quality)
    except AttributeError:
        pass  # very old PyMuPDF: fall back to structural compression only
    try:
        doc.subset_fonts()
    except Exception:
        pass
    doc.save(out, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    doc.close()
    if os.path.getsize(out) >= os.path.getsize(src):
        shutil.copyfile(src, out)  # never return a bigger file
    return out


# ---------------------------------------------------------------- PDF -> X

def pdf_to_word(src, out):
    from pdf2docx import Converter
    open_pdf(src).close()
    cv = Converter(str(src))
    try:
        cv.convert(str(out))
    finally:
        cv.close()
    return out


def pdf_to_powerpoint(src, out):
    from pptx import Presentation
    from pptx.util import Emu

    doc = open_pdf(src)
    prs = Presentation()
    first = doc[0].rect
    # PowerPoint slide limit is 56 inches (4032 pt)
    scale = min(1.0, 4032 / max(first.width, first.height))
    prs.slide_width = Emu(int(first.width * scale * 12700))
    prs.slide_height = Emu(int(first.height * scale * 12700))
    blank = prs.slide_layouts[6]
    for page in doc:
        pix = page.get_pixmap(dpi=150)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0, prs.slide_width, prs.slide_height)
        text = page.get_text("text").strip()
        if text:
            slide.notes_slide.notes_text_frame.text = text
    prs.save(out)
    return out


def pdf_to_excel(src, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    doc = open_pdf(src)
    wb = Workbook()
    wb.remove(wb.active)
    for page in doc:
        ws = wb.create_sheet(f"Page {page.number + 1}")
        row = 1
        tables = []
        try:
            tables = page.find_tables().tables
        except Exception:
            pass
        if tables:
            for t_index, table in enumerate(tables, 1):
                ws.cell(row=row, column=1, value=f"Table {t_index}").font = Font(bold=True)
                row += 1
                for data_row in table.extract():
                    for col, value in enumerate(data_row, 1):
                        ws.cell(row=row, column=col, value=_excel_value(value))
                    row += 1
                row += 1
        else:
            # No ruled tables: put each text line in a row, split on large gaps
            for block in page.get_text("blocks", sort=True):
                for line in block[4].splitlines():
                    if line.strip():
                        cells = [c for c in line.split("  ") if c.strip()]
                        for col, value in enumerate(cells, 1):
                            ws.cell(row=row, column=col, value=_excel_value(value.strip()))
                        row += 1
        for column in ws.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            ws.column_dimensions[column[0].column_letter].width = min(60, max(8, width + 2))
    wb.save(out)
    return out


def _excel_value(value):
    if value is None:
        return None
    text = str(value).strip()
    candidate = text.replace(",", "")
    try:
        return int(candidate) if candidate.lstrip("-").isdigit() else float(candidate)
    except ValueError:
        return text


def pdf_to_images(src, work, mode, dpi, fmt):
    doc = open_pdf(src)
    stem = clean_stem(src)
    outputs = []
    if mode == "extract":
        seen = set()
        for page in doc:
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                info = doc.extract_image(xref)
                if not info or info.get("width", 0) < 16:
                    continue
                out = work / f"{stem}_image_{len(outputs) + 1}.{info['ext']}"
                out.write_bytes(info["image"])
                outputs.append(out)
        if not outputs:
            raise ToolError("No embedded images were found in this PDF")
    else:
        ext = "png" if fmt == "png" else "jpg"
        for page in doc:
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            out = work / f"{stem}_page_{page.number + 1}.{ext}"
            if ext == "png":
                pix.save(out)
            else:
                out.write_bytes(pix.tobytes("jpg", jpg_quality=90))
            outputs.append(out)
    return outputs


# ---------------------------------------------------------------- X -> PDF

PAGE_SIZES = {"a4": (595.28, 841.89), "letter": (612, 792), "legal": (612, 1008)}


def images_to_pdf(files, out, page_size, orientation, margin):
    doc = fitz.open()
    for f in files:
        try:
            im = ImageOps.exif_transpose(Image.open(f))
        except Exception:
            raise ToolError(f"'{Path(f).name}' is not a valid image")
        if im.mode not in ("RGB", "L"):
            background = Image.new("RGB", im.size, "white")
            im = im.convert("RGBA")
            background.paste(im, mask=im.split()[-1])
            im = background
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=92)
        iw, ih = im.size

        if page_size == "fit":
            pw, ph = iw * 72 / 96, ih * 72 / 96
        else:
            pw, ph = PAGE_SIZES.get(page_size, PAGE_SIZES["a4"])
            landscape = orientation == "landscape" or (orientation == "auto" and iw > ih)
            if landscape:
                pw, ph = ph, pw
        page = doc.new_page(width=pw, height=ph)
        m = margin if page_size != "fit" else 0
        rect = fitz.Rect(m, m, pw - m, ph - m)
        page.insert_image(rect, stream=buf.getvalue(), keep_proportion=True)
    doc.save(out, garbage=3, deflate=True)
    return out


_office_lock = threading.Lock()


def office_to_pdf(src, out, kind):
    """Convert Word / PowerPoint / Excel using MS Office, falling back to LibreOffice."""
    src, out = str(Path(src).resolve()), str(Path(out).resolve())
    errors = []
    with _office_lock:
        try:
            _ms_office_to_pdf(src, out, kind)
            if os.path.exists(out):
                return out
        except Exception as e:  # noqa: BLE001 - try next engine
            errors.append(f"MS Office: {e}")
        soffice = _find_soffice()
        if soffice:
            outdir = tempfile.mkdtemp(prefix="lo_")
            try:
                subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, src],
                               check=True, timeout=300, capture_output=True)
                produced = Path(outdir) / (Path(src).stem + ".pdf")
                if produced.exists():
                    shutil.move(str(produced), out)
                    return out
            except Exception as e:  # noqa: BLE001
                errors.append(f"LibreOffice: {e}")
            finally:
                shutil.rmtree(outdir, ignore_errors=True)
    raise ToolError("Conversion failed. Microsoft Office or LibreOffice must be installed. " + " | ".join(errors))


def _ms_office_to_pdf(src, out, kind):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    try:
        if kind == "word":
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            try:
                d = app.Documents.Open(src, False, True, False)  # FileName, ConfirmConversions, ReadOnly, AddToRecent
                d.ExportAsFixedFormat(out, 17)  # wdExportFormatPDF
                d.Close(False)
            finally:
                app.Quit()
        elif kind == "powerpoint":
            app = win32com.client.DispatchEx("PowerPoint.Application")
            try:
                p = app.Presentations.Open(src, True, False, False)  # FileName, ReadOnly, Untitled, WithWindow
                p.SaveAs(out, 32)  # ppSaveAsPDF
                p.Close()
            finally:
                app.Quit()
        elif kind == "excel":
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            try:
                wb = app.Workbooks.Open(src, 0, True)  # FileName, UpdateLinks, ReadOnly
                wb.ExportAsFixedFormat(0, out)  # xlTypePDF
                wb.Close(False)
            finally:
                app.Quit()
    finally:
        pythoncom.CoUninitialize()


def _find_soffice():
    candidates = [
        shutil.which("soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    return next((c for c in candidates if c and os.path.exists(c)), None)


def _find_browser():
    pf = os.environ.get("ProgramFiles", r"C:\Program Files")
    pf86 = os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")
    local = os.environ.get("LOCALAPPDATA", "")
    candidates = [
        os.path.join(pf86, r"Microsoft\Edge\Application\msedge.exe"),
        os.path.join(pf, r"Microsoft\Edge\Application\msedge.exe"),
        os.path.join(pf, r"Google\Chrome\Application\chrome.exe"),
        os.path.join(pf86, r"Google\Chrome\Application\chrome.exe"),
        os.path.join(local, r"Google\Chrome\Application\chrome.exe"),
        shutil.which("msedge"), shutil.which("chrome"), shutil.which("chromium"),
    ]
    return next((c for c in candidates if c and os.path.exists(c)), None)


def html_to_pdf(out, url=None, html_file=None):
    if html_file:
        target = Path(html_file).resolve().as_uri()
    else:
        parsed = urlparse((url or "").strip())
        if parsed.scheme not in ("http", "https") or not parsed.netloc:
            raise ToolError("Enter a full web address starting with http:// or https://")
        target = parsed.geturl()
    browser = _find_browser()
    if not browser:
        raise ToolError("Microsoft Edge or Google Chrome is required for HTML to PDF")
    profile = tempfile.mkdtemp(prefix="html2pdf_")
    try:
        subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-first-run",
                        "--no-pdf-header-footer", f"--user-data-dir={profile}",
                        f"--print-to-pdf={Path(out).resolve()}", target],
                       timeout=180, capture_output=True)
    except subprocess.TimeoutExpired:
        raise ToolError("The page took too long to load")
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    if not os.path.exists(out) or os.path.getsize(out) == 0:
        raise ToolError("Could not render the page to PDF")
    return out


# ---------------------------------------------------------------- edit / sign / watermark

def pdf_preview(src, page_no, password=""):
    doc = open_pdf(src, password)
    page_no = max(0, min(page_no, doc.page_count - 1))
    page = doc[page_no]
    zoom = min(3.0, 900 / page.rect.width)
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    return {
        "page_count": doc.page_count,
        "page": page_no,
        "width": page.rect.width,
        "height": page.rect.height,
        "image": "data:image/png;base64," + base64.b64encode(pix.tobytes("png")).decode(),
        "lines": _page_lines(page),
    }


def render_pages(src, limit=30, width=900):
    doc = open_pdf(src)
    pages = []
    for page in doc.pages(0, min(limit, doc.page_count)):
        zoom = min(2.0, width / page.rect.width)
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        pages.append("data:image/jpeg;base64," + base64.b64encode(pix.tobytes("jpg", jpg_quality=82)).decode())
    return {"pages": pages, "total": doc.page_count}


def _page_lines(page):
    """Horizontal text lines in visible-page fractions, so the editor can offer them for editing."""
    W, H = page.rect.width, page.rect.height
    rot = page.rotation_matrix
    turn = fitz.Matrix(rot.a, rot.b, rot.c, rot.d, 0, 0)
    lines = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            d = fitz.Point(line["dir"]) * turn
            if abs(d.y) > 0.01 or d.x <= 0:
                continue
            r = fitz.Rect(line["bbox"]) * rot
            main = max(spans, key=lambda s: len(s["text"]))
            name, flags = main["font"].lower(), main["flags"]
            if flags & 8 or "courier" in name or "mono" in name:
                font = "courier"
            elif flags & 4 or "times" in name or "serif" in name and "sans" not in name:
                font = "times"
            else:
                font = "helvetica"
            lines.append({
                "x": r.x0 / W, "y": r.y0 / H, "w": r.width / W, "h": r.height / H,
                "baseline": (fitz.Point(spans[0]["origin"]) * rot).y / H,
                "text": "".join(s["text"] for s in line["spans"]).strip(),
                "size": round(main["size"], 1),
                "color": "#%06x" % main["color"],
                "font": font,
                "bold": bool(flags & 16) or "bold" in name,
            })
    return lines


FONTS = {"helvetica": "helv", "times": "tiro", "courier": "cour"}
BOLD_FONTS = {"helvetica": "hebo", "times": "tibo", "courier": "cobo"}


def edit_pdf(src, out, items_json, images):
    """Apply items placed in the browser editor. Coordinates are fractions (0..1) of the page."""
    try:
        items = json.loads(items_json or "[]")
    except json.JSONDecodeError:
        raise ToolError("Invalid edit data")
    if not items:
        raise ToolError("Add at least one element before saving")
    doc = open_pdf(src)
    prepared = set()

    # Existing text being replaced is removed first (real removal via redaction, not just covered up).
    replaced = {}
    for it in items:
        if it.get("type") == "replace":
            replaced.setdefault(int(it.get("page", 0)), []).append(it)
    for p, group in replaced.items():
        if not 0 <= p < doc.page_count:
            continue
        page = doc[p]
        page.remove_rotation()
        prepared.add(p)
        W, H = page.rect.width, page.rect.height
        for it in group:
            r = fitz.Rect(float(it["x"]) * W, float(it["y"]) * H,
                          (float(it["x"]) + float(it["w"])) * W, (float(it["y"]) + float(it["h"])) * H)
            inset = r.height * 0.15
            page.add_redact_annot(fitz.Rect(r.x0, r.y0 + inset, r.x1, r.y1 - inset), fill=False)
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,
                              graphics=getattr(fitz, "PDF_REDACT_LINE_ART_NONE", 0))

    for it in items:
        p = int(it.get("page", 0))
        if not 0 <= p < doc.page_count:
            continue
        page = doc[p]
        if p not in prepared:
            page.remove_rotation()  # makes visible coordinates == PDF coordinates
            prepared.add(p)
        W, H = page.rect.width, page.rect.height
        kind = it.get("type")
        if kind == "text":
            size = float(it.get("size", 14))
            x, y = float(it["x"]) * W, float(it["y"]) * H
            page.insert_text((x, y + size * 0.85), str(it.get("text", "")), fontsize=size,
                             fontname=FONTS.get(it.get("font"), "helv"), color=hex_to_rgb(it.get("color")))
        elif kind == "replace":
            text = str(it.get("text", "")).replace("\n", " ")
            if text.strip():
                fonts = BOLD_FONTS if it.get("bold") else FONTS
                page.insert_text((float(it["x"]) * W, float(it["baseline"]) * H), text,
                                 fontsize=float(it.get("size", 12)), fontname=fonts.get(it.get("font"), "helv"),
                                 color=hex_to_rgb(it.get("color")))
        elif kind in ("rect", "whiteout", "image"):
            rect = fitz.Rect(float(it["x"]) * W, float(it["y"]) * H,
                             (float(it["x"]) + float(it["w"])) * W, (float(it["y"]) + float(it["h"])) * H)
            if rect.is_empty:
                continue
            if kind == "whiteout":
                page.draw_rect(rect, color=None, fill=(1, 1, 1), width=0)
            elif kind == "rect":
                fill = hex_to_rgb(it["fill"]) if it.get("fill") else None
                page.draw_rect(rect, color=hex_to_rgb(it.get("color")), fill=fill,
                               width=float(it.get("stroke", 2)))
            else:
                data = images.get(it.get("image"))
                if data:
                    page.insert_image(rect, stream=data, keep_proportion=True, overlay=True)
        elif kind == "draw":
            points = [fitz.Point(float(px) * W, float(py) * H) for px, py in it.get("points", [])]
            if len(points) > 1:
                page.draw_polyline(points, color=hex_to_rgb(it.get("color")), width=float(it.get("stroke", 2)),
                                   lineCap=1, lineJoin=1)
    doc.save(out, garbage=3, deflate=True)
    return out


def watermark_pdf(src, out, form, image_bytes=None):
    doc = open_pdf(src)
    opacity = max(0.05, min(1.0, float(form.get("opacity", 40)) / 100))
    position = form.get("position", "center")
    pages = parse_pages(form.get("pages", "all"), doc.page_count)

    if image_bytes:
        im = Image.open(io.BytesIO(image_bytes)).convert("RGBA")
        alpha = im.split()[-1].point(lambda a: int(a * opacity))
        im.putalpha(alpha)
        buf = io.BytesIO()
        im.save(buf, "PNG")
        data = buf.getvalue()
        scale = float(form.get("image_scale", 40)) / 100
    else:
        text = (form.get("text") or "").strip()
        if not text:
            raise ToolError("Enter watermark text or upload an image")
        size = float(form.get("font_size", 48))
        color = hex_to_rgb(form.get("color", "#ff0000"))
        rotation = float(form.get("rotation", 45))
        font = FONTS.get(form.get("font"), "helv")

    for p in pages:
        page = doc[p]
        page.remove_rotation()
        r = page.rect
        if image_bytes:
            w = r.width * scale
            h = w * im.height / im.width
            cx, cy = _anchor(position, r, w, h)
            page.insert_image(fitz.Rect(cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), stream=data, overlay=True)
        else:
            tw = fitz.get_text_length(text, fontname=font, fontsize=size)
            cx, cy = _anchor(position, r, tw, size)
            origin = fitz.Point(cx - tw / 2, cy + size * 0.35)
            page.insert_text(origin, text, fontsize=size, fontname=font, color=color,
                             fill_opacity=opacity, stroke_opacity=opacity,
                             morph=(fitz.Point(cx, cy), fitz.Matrix(rotation)), overlay=True)
    doc.save(out, garbage=3, deflate=True)
    return out


def _anchor(position, r, w, h):
    pad = 30
    xs = {"left": pad + w / 2, "center": r.width / 2, "right": r.width - pad - w / 2}
    ys = {"top": pad + h / 2, "middle": r.height / 2, "bottom": r.height - pad - h / 2}
    if position == "center":
        return xs["center"], ys["middle"]
    v, _, hz = position.partition("-")
    return xs.get(hz, xs["center"]), ys.get(v, ys["middle"])


def add_page_numbers(src, out, position, start, size):
    doc = open_pdf(src)
    for i, page in enumerate(doc):
        page.remove_rotation()
        label = str(start + i)
        tw = fitz.get_text_length(label, fontname="helv", fontsize=size)
        r = page.rect
        x = {"left": 36, "center": (r.width - tw) / 2, "right": r.width - 36 - tw}[position.split("-")[1]]
        y = 30 + size if position.startswith("top") else r.height - 24
        page.insert_text((x, y), label, fontsize=size, fontname="helv", color=(0, 0, 0))
    doc.save(out, garbage=3, deflate=True)
    return out


# ---------------------------------------------------------------- security

def unlock_pdf(src, out, password):
    try:
        doc = fitz.open(src)
    except Exception:
        raise ToolError("Not a valid PDF file")
    if doc.needs_pass and not doc.authenticate(password or ""):
        raise ToolError("Wrong password")
    doc.save(out, encryption=fitz.PDF_ENCRYPT_NONE, garbage=3, deflate=True)
    return out


def protect_pdf(src, out, password):
    if not password or len(password) < 4:
        raise ToolError("Password must be at least 4 characters")
    doc = open_pdf(src)
    perms = fitz.PDF_PERM_PRINT | fitz.PDF_PERM_COPY | fitz.PDF_PERM_ACCESSIBILITY
    doc.save(out, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password,
             permissions=perms, garbage=3, deflate=True)
    return out


# ---------------------------------------------------------------- images

IMAGE_FORMATS = {"jpeg": ("JPEG", "jpg"), "png": ("PNG", "png"), "webp": ("WEBP", "webp")}


def _load_image(path):
    try:
        im = Image.open(path)
        im.load()
    except Exception:
        raise ToolError(f"'{Path(path).name}' is not a valid image")
    return ImageOps.exif_transpose(im)


def _output_format(im, requested):
    if requested in IMAGE_FORMATS:
        return requested
    fmt = (im.format or "").lower()
    return {"jpg": "jpeg", "jpeg": "jpeg", "png": "png", "webp": "webp"}.get(fmt, "jpeg")


def _encode(im, fmt, quality, dpi=None):
    pil_fmt, _ = IMAGE_FORMATS[fmt]
    if fmt == "jpeg" and im.mode not in ("RGB", "L"):
        bg = Image.new("RGB", im.size, "white")
        rgba = im.convert("RGBA")
        bg.paste(rgba, mask=rgba.split()[-1])
        im = bg
    elif fmt == "webp" and im.mode not in ("RGB", "RGBA"):
        im = im.convert("RGBA")
    buf = io.BytesIO()
    kwargs = {"optimize": True}
    if dpi:
        kwargs["dpi"] = (dpi, dpi)
    if fmt in ("jpeg", "webp"):
        kwargs["quality"] = int(quality)
        if fmt == "jpeg":
            kwargs["progressive"] = True
    elif fmt == "png" and quality < 90:
        im = im.convert("RGBA").quantize(colors=256, method=Image.Quantize.FASTOCTREE) \
            if im.mode in ("RGBA", "LA", "P") else im.convert("RGB").quantize(colors=256)
        kwargs = {"optimize": True}
    im.save(buf, pil_fmt, **kwargs)
    return buf.getvalue()


def _encode_to_target(im, fmt, target_bytes, dpi=None):
    """Find the best quality (and if needed, smaller dimensions) that fits target_bytes."""
    current = im
    for _ in range(15):
        if fmt == "png":
            data = _encode(current, fmt, 50, dpi)
            if len(data) <= target_bytes:
                return data
        else:
            lo, hi, best = 5, 95, None
            while lo <= hi:
                q = (lo + hi) // 2
                data = _encode(current, fmt, q, dpi)
                if len(data) <= target_bytes:
                    best, lo = data, q + 1
                else:
                    hi = q - 1
            if best:
                return best
        w, h = current.size
        if w < 32 or h < 32:
            break
        current = current.resize((max(1, int(w * 0.85)), max(1, int(h * 0.85))), Image.LANCZOS)
    raise ToolError(f"Could not reach {target_bytes // 1024} KB. Try a larger target size.")


def compress_image(src, work, quality, target_kb, fmt):
    im = _load_image(src)
    fmt = _output_format(im, fmt)
    if target_kb > 0:
        data = _encode_to_target(im, fmt, target_kb * 1024)
    else:
        data = _encode(im, fmt, quality)
    out = work / f"{clean_stem(src)}_compressed.{IMAGE_FORMATS[fmt][1]}"
    out.write_bytes(data)
    return out


UNIT_TO_INCH = {"in": 1, "cm": 1 / 2.54, "mm": 1 / 25.4}


def resize_image(src, work, form):
    im = _load_image(src)
    ow, oh = im.size
    unit = form.get("unit", "px")
    dpi = int(form.get("dpi") or 300)
    fit = form.get("fit", "contain")

    def to_px(value):
        if value in (None, ""):
            return None
        v = float(value)
        if v <= 0:
            raise ToolError("Width and height must be greater than 0")
        if unit == "percent":
            return v
        if unit in UNIT_TO_INCH:
            return round(v * UNIT_TO_INCH[unit] * dpi)
        return round(v)

    if unit == "percent":
        pct = to_px(form.get("percent"))
        if not pct:
            raise ToolError("Enter a percentage")
        w, h = max(1, round(ow * pct / 100)), max(1, round(oh * pct / 100))
        fit = "stretch"
    else:
        w, h = to_px(form.get("width")), to_px(form.get("height"))
    if not w and not h:
        raise ToolError("Enter a width and/or a height")
    if not w:
        w = round(ow * h / oh)
    if not h:
        h = round(oh * w / ow)
    if w > 20000 or h > 20000:
        raise ToolError("Maximum size is 20000 x 20000 pixels")

    if fit == "stretch":
        result = im.resize((w, h), Image.LANCZOS)
    elif fit == "cover":
        result = ImageOps.fit(im, (w, h), Image.LANCZOS)
    elif fit == "pad":
        bg = form.get("background", "#ffffff")
        result = ImageOps.pad(im.convert("RGBA") if im.mode == "P" else im, (w, h), Image.LANCZOS,
                              color=bg if im.mode != "RGBA" else None)
    else:  # contain: keep ratio inside the box
        result = ImageOps.contain(im, (w, h), Image.LANCZOS)

    fmt = _output_format(im, form.get("format", "same"))
    target_kb = int(form.get("target_kb") or 0)
    if target_kb > 0:
        data = _encode_to_target(result, fmt, target_kb * 1024, dpi)
    else:
        data = _encode(result, fmt, int(form.get("quality") or 92), dpi)
    out = work / f"{clean_stem(src)}_{result.width}x{result.height}.{IMAGE_FORMATS[fmt][1]}"
    out.write_bytes(data)
    return out


def convert_image(src, work, fmt):
    im = _load_image(src)
    if fmt not in IMAGE_FORMATS:
        raise ToolError("Choose an output format")
    out = work / f"{clean_stem(src)}.{IMAGE_FORMATS[fmt][1]}"
    out.write_bytes(_encode(im, fmt, 95))
    return out


REMBG_MODELS = {"general": "isnet-general-use", "person": "u2net_human_seg", "fast": "u2netp"}
_rembg_sessions = {}
_rembg_lock = threading.Lock()


def _plain_edge_box(im, alpha):
    """Crop box excluding edge strips that are one flat color yet kept as subject (photo/scan borders)."""
    import numpy as np

    rgb = np.asarray(im.convert("RGB"), dtype=np.int16)
    fg = np.asarray(alpha) > 128
    h, w = fg.shape

    def strip(pixels, mask, limit):
        n = 0
        for px, m in zip(pixels, mask):
            if n >= limit or m.mean() < 0.97 or px.std(axis=0).max() > 8:
                break
            n += 1
        return n if n >= 2 else 0

    cols, cols_fg = rgb.transpose(1, 0, 2), fg.T
    top = strip(rgb, fg, h // 4)
    bottom = strip(rgb[::-1], fg[::-1], h // 4)
    left = strip(cols, cols_fg, w // 4)
    right = strip(cols[::-1], cols_fg[::-1], w // 4)
    return left, top, w - right, h - bottom


def _clean_edges(im, cut, strength):
    """Remove the halo of the old background: shrink + feather the mask, then un-mix the old colour from edge pixels."""
    import numpy as np

    if strength == "off":
        return cut
    w, h = cut.size
    shrink = max(1, min(6, round(min(w, h) / 350))) * (2 if strength == "strong" else 1)
    alpha = cut.getchannel("A").filter(ImageFilter.MinFilter(2 * shrink + 1))
    alpha = alpha.filter(ImageFilter.GaussianBlur(max(0.6, shrink * 0.6)))

    a = np.asarray(alpha, dtype=np.float32) / 255
    rgb = np.asarray(im.convert("RGB"), dtype=np.float32)
    from scipy.ndimage import distance_transform_edt

    # Plain old background: key out its colour in the outer band near dark detail (hair gaps the AI kept).
    model_alpha = np.asarray(cut.getchannel("A"))
    old_bg = rgb[model_alpha < 10]
    if len(old_bg) > 500 and old_bg.std(axis=0).max() < 25:
        bg_color = np.median(old_bg, axis=0)
        band_px = max(6, round(min(w, h) * (0.08 if strength == "strong" else 0.05)))
        inside = a > 0.01
        band = inside & (distance_transform_edt(inside) <= band_px)
        dark = rgb.mean(axis=2) < 110
        near_dark = distance_transform_edt(~dark) <= band_px / 2 if dark.any() else np.zeros_like(dark)
        # Light areas reaching deep inside the subject (e.g. a white shirt) are subject, not background gaps.
        from scipy.ndimage import label
        chroma_all = rgb.max(axis=2) - rgb.min(axis=2)
        light = inside & (chroma_all <= 30) & (rgb.mean(axis=2) > float(bg_color.mean()) - 115)
        labels, _ = label(light)
        deep_ids = np.unique(labels[light & (distance_transform_edt(inside) > band_px)])
        deep_light = np.isin(labels, deep_ids[deep_ids > 0])
        key = band & near_dark & ~deep_light
        dist = np.linalg.norm(rgb[key] - bg_color, axis=1)
        key_alpha = np.clip((dist - 15) / 45, 0, 1)
        lum = rgb.mean(axis=2)
        bg_lum = float(bg_color.mean())
        hair_lum = float(np.percentile(lum[key & dark], 20)) if (key & dark).any() else 0.0
        if bg_lum - hair_lum > 60:
            # Dark detail on a light background: opacity follows how dark the pixel is (mixed hair/background pixels).
            lum_alpha = np.clip((bg_lum - 12 - lum[key]) / (bg_lum - 12 - hair_lum - 25), 0, 1)
            chroma = rgb[key].max(axis=1) - rgb[key].min(axis=1)
            lum_alpha[chroma > 30] = 1.0  # coloured pixels (skin) are subject, not hair/background mix
            key_alpha = np.minimum(key_alpha, lum_alpha)
        a[key] = np.minimum(a[key], key_alpha)

    solid = a >= 0.98
    edge = (a > 0.01) & ~solid
    if solid.any() and edge.any():
        # Edge pixels take the colour of the nearest solid subject pixel, so no old-background colour leaks in.
        _, (iy, ix) = distance_transform_edt(~solid, return_indices=True)
        rgb[edge] = rgb[iy[edge], ix[edge]]
    out = np.dstack([rgb, a * 255]).round().astype(np.uint8)
    return Image.fromarray(out, "RGBA")


def remove_background(src, work, form, bg_image_bytes=None):
    try:
        from rembg import new_session, remove
    except ImportError:
        raise ToolError("Background removal needs the 'rembg' package. Run run.bat again to install it.")

    im = _load_image(src).convert("RGB")
    model = REMBG_MODELS.get(form.get("model"), REMBG_MODELS["general"])
    with _rembg_lock:
        if model not in _rembg_sessions:
            try:
                _rembg_sessions[model] = new_session(model)  # downloads the AI model on first use
            except Exception as e:
                raise ToolError(f"Could not load the AI model (internet needed on first use): {e}")
        cut = remove(im, session=_rembg_sessions[model], post_process_mask=False, decontaminate=True)

    if form.get("trim", "yes") == "yes":
        box = _plain_edge_box(im, cut.getchannel("A"))
        if box != (0, 0, im.width, im.height):
            im, cut = im.crop(box), cut.crop(box)
    mode = form.get("background", "transparent")
    if mode == "color":
        bg = Image.new("RGBA", cut.size, form.get("color") or "#ffffff")
    elif mode == "image":
        if not bg_image_bytes:
            raise ToolError("Upload a background image")
        try:
            bg = Image.open(io.BytesIO(bg_image_bytes))
            bg = ImageOps.fit(ImageOps.exif_transpose(bg).convert("RGBA"), cut.size, Image.LANCZOS)
        except Exception:
            raise ToolError("The background image is not valid")
    elif mode == "blur":
        radius = max(2, min(60, int(float(form.get("blur") or 18))))
        bg = im.filter(ImageFilter.GaussianBlur(radius * max(im.size) / 1000)).convert("RGBA")
    else:
        bg = None
    result = cut
    if bg is not None:
        bg.alpha_composite(cut)
        result = bg

    fmt = form.get("format", "png")
    fmt = fmt if fmt in IMAGE_FORMATS else "png"
    out = work / f"{clean_stem(src)}_{'no_bg' if bg is None else 'new_bg'}.{IMAGE_FORMATS[fmt][1]}"
    out.write_bytes(_encode(result, fmt, 95))
    return out
